"""Build the idle needle import from the supplied workbook without editing it."""

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Excel" / "IdleNeedleUnitPrice.xlsx"
OUTPUT = ROOT / "supabase" / "migrations" / "20261023000000_idle_needle_master.sql"


def sql_text(value):
    return "'" + str(value).strip().replace("'", "''") + "'"


def build_sql():
    values_book = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
    formulas_book = openpyxl.load_workbook(SOURCE, data_only=False, read_only=True)
    sheet = values_book["UNIT PRICE"]
    formulas = formulas_book["UNIT PRICE"]
    assert list(next(sheet.values)) == [
        "DRISCRIPTION ITEM", "ITEM CODE", "NEEDLE TYPE", "SIZE/DIA", "BRAND",
        "UNIT PRICE (BDT)", "UNIT PRICE (USD)",
    ], "Unexpected workbook columns"
    rows = []
    codes = set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value is not None for value in row):
            continue
        assert all(value is not None and str(value).strip() for value in row), (
            f"Missing data at worksheet row {row_number}"
        )
        description, code, needle_type, size, brand, bdt, usd = row
        code = str(code).strip()
        assert code not in codes, f"Duplicate item code: {code}"
        codes.add(code)
        bdt, usd = Decimal(str(bdt)), Decimal(str(usd))
        assert bdt.is_finite() and usd.is_finite() and bdt >= 0 and usd >= 0
        # Verify cached USD values against the workbook's actual formula.
        assert formulas.cell(row_number, 7).value == (
            "=SUM(PRICE[[#This Row],[UNIT PRICE (BDT)]]/110)"
        ), f"Unexpected USD formula at G{row_number}"
        assert abs(usd - bdt / Decimal(110)) < Decimal("0.000000000001")
        bdt_sql = bdt.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
        usd_sql = usd.quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
        fields = [sql_text(v) for v in [description, code, needle_type, size, brand]]
        fields += [str(bdt_sql), str(usd_sql), "true"]
        rows.append("    (" + ", ".join(fields) + ")")
    values_book.close()
    formulas_book.close()
    assert codes == {f"IN-{number:03d}" for number in range(1, 42)}, "Unexpected item coverage"

    return """-- =============================================================================
-- Migration: 20261023000000_idle_needle_master.sql
-- Source: Excel/IdleNeedleUnitPrice.xlsx, 'UNIT PRICE'!A2:G42 (41 items).
-- Run after 20241001000008_needle_master.sql on an existing ERP database.
-- Imports IN-001 through IN-041 and selects every currently active floor.
-- The existing application calls the idle needle group 'ideal_needle'.
-- Prices use the workbook's BDT and USD values, rounded to the database scales
-- (4 and 6 decimals respectively). The workbook's USD conversion is BDT / 110.
-- Safe to rerun: updates these items and adds missing idle floor assignments.
-- Preserves other items, TAC codes, and all existing floor/group assignments.
-- Future floors can be selected by rerunning this import or through the app.
-- =============================================================================

BEGIN;

DO $$
BEGIN
  IF to_regclass('public.needle_items') IS NULL
     OR to_regclass('public.needle_floor_assignments') IS NULL
     OR to_regclass('public.floors') IS NULL THEN
    RAISE EXCEPTION 'Run the ERP schema and 20241001000008_needle_master.sql first.';
  END IF;
  IF NOT EXISTS (SELECT 1 FROM public.floors WHERE active = true) THEN
    RAISE EXCEPTION 'Create or enable at least one floor before importing idle needles.';
  END IF;
END;
$$;

WITH imported AS (
  INSERT INTO public.needle_items (
    description_item, item_code, needle_type, size_dia, brand,
    unit_price_bdt, unit_price_usd, active
  ) VALUES
""" + ",\n".join(rows) + """
  ON CONFLICT (item_code) DO UPDATE SET
    description_item = EXCLUDED.description_item,
    needle_type = EXCLUDED.needle_type,
    size_dia = EXCLUDED.size_dia,
    brand = EXCLUDED.brand,
    unit_price_bdt = EXCLUDED.unit_price_bdt,
    unit_price_usd = EXCLUDED.unit_price_usd,
    active = true
  RETURNING id
), assigned AS (
  INSERT INTO public.needle_floor_assignments (
    needle_item_id, floor_id, needle_group
  )
  SELECT i.id, f.id, 'ideal_needle'
  FROM imported AS i
  CROSS JOIN public.floors AS f
  WHERE f.active = true
  ON CONFLICT (needle_item_id, floor_id, needle_group) DO NOTHING
  RETURNING needle_item_id
)
SELECT
  (SELECT count(*) FROM imported) AS imported_idle_needles,
  (SELECT count(*) FROM public.floors WHERE active = true) AS selected_active_floors,
  (SELECT count(*) FROM assigned) AS added_floor_assignments;

NOTIFY pgrst, 'reload schema';

COMMIT;
"""


if __name__ == "__main__":
    sql = build_sql()
    OUTPUT.write_text(sql, encoding="utf-8")
    print(f"Created {OUTPUT.name}: 41 validated items with all active idle floors selected.")
