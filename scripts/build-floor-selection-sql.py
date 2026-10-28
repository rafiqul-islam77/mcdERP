"""Generate the floor assignment SQL from the supplied workbook (read-only)."""
from collections import Counter
from pathlib import Path
import hashlib
import argparse
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--source', default='C:/Users/tazzi/Downloads/FLOOR WISE ITEM CODE.xlsx')
parser.add_argument('--output', default='supabase/migrations/20261027010000_floor_selections_from_workbook.sql')
parser.add_argument('--date', default='2026-10-27')
args = parser.parse_args()
SOURCE = Path(args.source)
OUTPUT = ROOT / args.output
workbook = openpyxl.load_workbook(SOURCE, data_only=False)
stock_format = workbook.sheetnames == ['RUNNING+NEW', 'IDLE']
assert stock_format or workbook.sheetnames == ['NEW', 'RUNNING', 'IDLE']
headers = list('ABCDE') + ['EMB'] + list('FGHIJK') + ['SAMPLE', 'TRAINING CENTRE']
floor_map = {name: 'FL-' + name for name in headers}
floor_map.update({'SAMPLE': 'FL-SMP', 'TRAINING CENTRE': 'FL-TC'})
groups = {(category, code): [] for category in ('new', 'running', 'idle') for code in floor_map.values()}
source_cells = []

def resolve_floor(label):
    label = str(label).strip().upper()
    if label.startswith('SAMPLE-'):
        return 'FL-SMP'
    if label.startswith('TRAINING CENTRE'):
        return 'FL-TC'
    if label.startswith('EMB-'):
        return 'FL-EMB'
    match = re.match(r'^([A-K])-FLOOR\s*\(', label)
    assert match, label
    return 'FL-' + match[1]

if stock_format:
    assert [c.value for c in workbook['RUNNING+NEW'][1]][:3] == ['FLOOR NAME & LOCATION', 'CATEGORY', 'ITEM CODE']
    assert [c.value for c in workbook['IDLE'][1]][:2] == ['FLOOR NAME & LOCATION', 'ITEM CODE']
    for sheet in workbook:
        for row in sheet.iter_rows(min_row=2):
            values = [cell.value for cell in row]
            label = values[0]
            code = values[2] if sheet.title == 'RUNNING+NEW' else values[1]
            if code is None:
                # Only explicitly identified subtotal/grand-total rows may be skipped.
                assert (sheet.title == 'RUNNING+NEW' and (values[1] in ('NEW Total', 'RUNNING Total') or label == 'Grand Total')) or (sheet.title == 'IDLE' and str(label).endswith(' Total')), (sheet.title, row[0].row, values)
                continue
            category = str(values[1]).lower() if sheet.title == 'RUNNING+NEW' else 'idle'
            assert category in ('new', 'running', 'idle'), values
            code = str(code).strip().upper()
            assert re.fullmatch(r'IN-\d{3}' if category == 'idle' else r'SN-\d{3}', code), values
            floor = resolve_floor(label)
            # Selection is based on item presence, regardless of stock quantity.
            groups[(category, floor)].append(code)
            source_cells.append((category, floor, code))

for title in (() if stock_format else ('NEW', 'RUNNING')):
    sheet = workbook[title]
    assert [cell.value for cell in sheet[2]] == headers
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            if cell.value is None:
                continue
            code = str(cell.value).strip().upper()
            assert re.fullmatch(r'SN-\d{3}', code), (title, cell.coordinate, code)
            floor = floor_map[headers[cell.column - 1]]
            groups[(title.lower(), floor)].append(code)
            source_cells.append((title.lower(), floor, code))
for row in ([] if stock_format else workbook['IDLE'].iter_rows()):
    values = [cell.value for cell in row]
    if not any(value is not None for value in values):
        continue
    assert values[0] and values[1] and not any(value is not None for value in values[2:])
    label, code = str(values[0]).strip(), str(values[1]).strip().upper()
    assert re.fullmatch(r'IN-\d{3}', code), (row[1].coordinate, code)
    if label.startswith('SAMPLE-'):
        floor = 'FL-SMP'
    elif label.startswith('TRAINING CENTRE('):
        floor = 'FL-TC'
    else:
        match = re.match(r'^([A-K])-FLOOR\s*\(', label)
        assert match, label
        floor = 'FL-' + match[1]
    groups[('idle', floor)].append(code)
    source_cells.append(('idle', floor, code))

assert not groups[('idle', 'FL-EMB')]  # User explicitly requested clearing IDLE EMB.
assert len(source_cells) == len(set(source_cells)), 'Duplicate item/floor/category in source'
seed = (ROOT / 'supabase/migrations/20261023000002_reset_separate_needle_catalogs.sql').read_text()
known = set(re.findall(r'\b(?:SN|IN)-\d+', seed))
assert {row[2] for row in source_cells} <= known, 'Unknown seeded item code'
counts = Counter(row[0] for row in source_cells)

sql = f"""-- Source: {SOURCE.name}, sheets {', '.join(workbook.sheetnames)}.
-- Source SHA256: {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}
-- Generated {args.date}. Run the ENTIRE file in Supabase SQL Editor.
-- Select all listed item codes regardless of stock quantity; exclude subtotal rows.
-- Requires the separate normal/idle catalogs and entry_type migration
-- (20261023000002_reset_separate_needle_catalogs.sql and
--  20261024000000_needle_floor_entry_type.sql) to have already been applied.
-- Exact replacement on the 14 mapped floors: unlisted item selections are removed.
-- IDLE EMB is cleared explicitly at the user's request because it is absent in IDLE.
-- Floors outside this mapping, item master data and daily transactions are unchanged.
-- Expected selections: NEW {counts['new']}, RUNNING {counts['running']}, IDLE {counts['idle']}.
-- Atomic and repeat-safe. Missing database item/floor codes abort before mutations.

BEGIN;

-- Keep the validated mapping stable and prevent concurrent assignment edits.
LOCK TABLE public.floors, public.normal_needle_items, public.idle_needle_items IN SHARE MODE;
LOCK TABLE public.normal_needle_floor_assignments,
           public.idle_needle_floor_assignments IN SHARE ROW EXCLUSIVE MODE;

CREATE TEMP TABLE floor_selection_source (
  category text NOT NULL CHECK (category IN ('new', 'running', 'idle')),
  floor_code text NOT NULL,
  item_codes text[] NOT NULL,
  PRIMARY KEY (category, floor_code)
) ON COMMIT DROP;

INSERT INTO floor_selection_source (category, floor_code, item_codes) VALUES
"""
values = []
for (category, floor), codes in groups.items():
    array = 'ARRAY[' + ', '.join("'" + code + "'" for code in codes) + ']::text[]'
    values.append(f"  ('{category}', '{floor}', {array})")
sql += ',\n'.join(values) + ';\n'
sql += """
CREATE TEMP TABLE desired_floor_selections ON COMMIT DROP AS
SELECT s.category, s.floor_code, code.item_code
FROM floor_selection_source AS s
CROSS JOIN LATERAL unnest(s.item_codes) AS code(item_code);

-- Validate exact database matches before deleting any existing selections.
DO $$
DECLARE missing text;
BEGIN
  SELECT string_agg(s.floor_code, ', ' ORDER BY s.floor_code) INTO missing
  FROM (SELECT DISTINCT floor_code FROM floor_selection_source) AS s
  WHERE NOT EXISTS (SELECT 1 FROM public.floors f WHERE f.code = s.floor_code);
  IF missing IS NOT NULL THEN
    RAISE EXCEPTION 'Missing floor codes: %. No selections changed.', missing;
  END IF;

  SELECT string_agg(s.item_code, ', ' ORDER BY s.item_code) INTO missing
  FROM (SELECT DISTINCT item_code FROM desired_floor_selections WHERE category IN ('new', 'running')) AS s
  WHERE NOT EXISTS (SELECT 1 FROM public.normal_needle_items n WHERE n.item_code = s.item_code);
  IF missing IS NOT NULL THEN
    RAISE EXCEPTION 'Missing normal needle item codes: %. No selections changed.', missing;
  END IF;

  SELECT string_agg(s.item_code, ', ' ORDER BY s.item_code) INTO missing
  FROM (SELECT DISTINCT item_code FROM desired_floor_selections WHERE category = 'idle') AS s
  WHERE NOT EXISTS (SELECT 1 FROM public.idle_needle_items n WHERE n.item_code = s.item_code);
  IF missing IS NOT NULL THEN
    RAISE EXCEPTION 'Missing idle needle item codes: %. No selections changed.', missing;
  END IF;
END $$;

CREATE TEMP TABLE resolved_floor_selections ON COMMIT DROP AS
SELECT s.category, f.id AS floor_id, n.id AS needle_item_id
FROM desired_floor_selections s
JOIN public.floors f ON f.code = s.floor_code
JOIN public.normal_needle_items n ON n.item_code = s.item_code
WHERE s.category IN ('new', 'running')
UNION ALL
SELECT s.category, f.id, n.id
FROM desired_floor_selections s
JOIN public.floors f ON f.code = s.floor_code
JOIN public.idle_needle_items n ON n.item_code = s.item_code
WHERE s.category = 'idle';

-- Remove only assignments absent from the supplied list on the mapped floors.
-- Existing desired assignments retain their IDs and timestamps.
DELETE FROM public.normal_needle_floor_assignments a
USING public.floors f, floor_selection_source s
WHERE a.floor_id = f.id AND f.code = s.floor_code
  AND s.category = a.entry_type AND s.category IN ('new', 'running')
  AND NOT EXISTS (
    SELECT 1 FROM resolved_floor_selections d
    WHERE d.category = a.entry_type AND d.floor_id = a.floor_id
      AND d.needle_item_id = a.needle_item_id
  );

DELETE FROM public.idle_needle_floor_assignments a
USING public.floors f, floor_selection_source s
WHERE a.floor_id = f.id AND f.code = s.floor_code AND s.category = 'idle'
  AND NOT EXISTS (
    SELECT 1 FROM resolved_floor_selections d
    WHERE d.category = 'idle' AND d.floor_id = a.floor_id
      AND d.needle_item_id = a.needle_item_id
  );

INSERT INTO public.normal_needle_floor_assignments (needle_item_id, floor_id, entry_type)
SELECT needle_item_id, floor_id, category
FROM resolved_floor_selections WHERE category IN ('new', 'running')
ON CONFLICT (needle_item_id, floor_id, entry_type) DO NOTHING;

INSERT INTO public.idle_needle_floor_assignments (needle_item_id, floor_id)
SELECT needle_item_id, floor_id FROM resolved_floor_selections WHERE category = 'idle'
ON CONFLICT (needle_item_id, floor_id) DO NOTHING;

CREATE TEMP TABLE actual_floor_selections ON COMMIT DROP AS
SELECT a.entry_type AS category, a.floor_id, a.needle_item_id
FROM public.normal_needle_floor_assignments a
JOIN public.floors f ON f.id = a.floor_id
JOIN floor_selection_source s ON s.floor_code = f.code AND s.category = a.entry_type
UNION ALL
SELECT 'idle', a.floor_id, a.needle_item_id
FROM public.idle_needle_floor_assignments a
JOIN public.floors f ON f.id = a.floor_id
JOIN floor_selection_source s ON s.floor_code = f.code AND s.category = 'idle';

-- Roll back the entire script if either missing or extra selections remain.
DO $$
BEGIN
  IF EXISTS (SELECT * FROM resolved_floor_selections EXCEPT SELECT * FROM actual_floor_selections)
     OR EXISTS (SELECT * FROM actual_floor_selections EXCEPT SELECT * FROM resolved_floor_selections) THEN
    RAISE EXCEPTION 'Floor selection verification failed. All changes rolled back.';
  END IF;
END $$;

-- Result grid: expected_count and actual_count must match for every row.
SELECT upper(s.category) AS category, s.floor_code,
       cardinality(s.item_codes) AS expected_count,
       count(a.needle_item_id) AS actual_count
FROM floor_selection_source s
JOIN public.floors f ON f.code = s.floor_code
LEFT JOIN actual_floor_selections a ON a.category = s.category AND a.floor_id = f.id
GROUP BY s.category, s.floor_code, s.item_codes
ORDER BY CASE s.category WHEN 'new' THEN 1 WHEN 'running' THEN 2 ELSE 3 END, s.floor_code;

COMMIT;
"""
OUTPUT.write_text(sql, encoding='utf-8')

# Re-read generated values independently and reconcile every item/floor/category.
generated = OUTPUT.read_text(encoding='utf-8')
parsed = re.findall(r"\('([^']+)', '([^']+)', ARRAY\[(.*?)\]::text\[\]\)", generated)
actual = [(category, floor, code) for category, floor, codes in parsed
          for code in re.findall(r"'([^']+)'", codes)]
assert len(parsed) == 42
assert Counter(actual) == Counter(source_cells)
assert sum(category == 'idle' and floor == 'FL-EMB' and codes == '' for category, floor, codes in parsed) == 1
print({'output': str(OUTPUT), 'assignments': dict(counts), 'floor_category_pairs': len(parsed),
       'unique_normal_items': len({code for category, _, code in actual if category != 'idle'}),
       'unique_idle_items': len({code for category, _, code in actual if category == 'idle'}),
       'verification': 'Every generated mapping matches the source; IDLE EMB is explicitly empty.'})
